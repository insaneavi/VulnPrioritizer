from io import BytesIO
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY="172033"; BLUE="2563EB"; WHITE="FFFFFF"; RED="FECACA"; RED_DARK="991B1B"
ORANGE="FED7AA"; ORANGE_DARK="9A3412"; YELLOW="FEF3C7"; GREEN="DCFCE7"; GRAY="E5E7EB"; LIGHT="F8FAFC"
THIN=Side(style="thin",color="D1D5DB")

def _style_header(ws,row=1):
    for c in ws[row]:
        c.fill=PatternFill("solid",fgColor=NAVY); c.font=Font(color=WHITE,bold=True); c.alignment=Alignment(vertical="center")
    ws.row_dimensions[row].height=24

def _finish(ws, freeze="A2", filter_range=None):
    ws.freeze_panes=freeze
    if filter_range: ws.auto_filter.ref=filter_range
    for col in range(1,ws.max_column+1):
        vals=[str(ws.cell(r,col).value or "") for r in range(1,min(ws.max_row,250)+1)]
        width=min(max(10,max((len(x) for x in vals),default=8)+2),42)
        ws.column_dimensions[get_column_letter(col)].width=width
    for row in ws.iter_rows():
        for c in row:
            c.border=Border(bottom=THIN); c.alignment=Alignment(vertical="top",wrap_text=True)

def _table(ws,name):
    if ws.max_row<2: return
    ref=f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    t=Table(displayName=name,ref=ref); t.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True,showFirstColumn=False,showLastColumn=False)
    ws.add_table(t)

def _priority_reason(v):
    reasons=[]
    if v.get("kev"): reasons.append("KNOWN EXPLOITED (KEV)")
    if (v.get("epss_raw") or 0)>=.10: reasons.append("HIGH EPSS")
    if (v.get("cvss_raw") or 0)>=9: reasons.append("CRITICAL CVSS")
    return " + ".join(reasons)

def _signal_colors(ws, headers):
    if ws.max_row < 2:
        return
    idx={c.value:i+1 for i,c in enumerate(ws[1])}
    if "CISA KEV" in idx:
        col=get_column_letter(idx["CISA KEV"]); ws.conditional_formatting.add(f"{col}2:{col}{ws.max_row}",CellIsRule(operator="equal",formula=['"YES"'],fill=PatternFill("solid",fgColor=RED),font=Font(color=RED_DARK,bold=True)))
    if "EPSS" in idx:
        col=get_column_letter(idx["EPSS"]); rng=f"{col}2:{col}{ws.max_row}"
        ws.conditional_formatting.add(rng,CellIsRule(operator="greaterThanOrEqual",formula=["0.5"],fill=PatternFill("solid",fgColor=ORANGE),font=Font(color=ORANGE_DARK,bold=True)))
        ws.conditional_formatting.add(rng,CellIsRule(operator="between",formula=["0.1","0.499999"],fill=PatternFill("solid",fgColor=YELLOW)))
    if "CVSS" in idx:
        col=get_column_letter(idx["CVSS"]); rng=f"{col}2:{col}{ws.max_row}"
        ws.conditional_formatting.add(rng,CellIsRule(operator="greaterThanOrEqual",formula=["9"],fill=PatternFill("solid",fgColor=RED),font=Font(color=RED_DARK,bold=True)))
        ws.conditional_formatting.add(rng,CellIsRule(operator="between",formula=["7","8.999"],fill=PatternFill("solid",fgColor=ORANGE)))
        ws.conditional_formatting.add(rng,CellIsRule(operator="between",formula=["4","6.999"],fill=PatternFill("solid",fgColor=YELLOW)))
        ws.conditional_formatting.add(rng,CellIsRule(operator="lessThan",formula=["4"],fill=PatternFill("solid",fgColor=GREEN)))

def _write_vulns(ws, rows):
    headers=["Attention Reason","CVE","Title","CISA KEV","EPSS","EPSS Percentile","CVSS","Affected Assets","Total Findings","Published Date","Age (Days)","Rapid7 Vulnerability ID"]
    ws.append(headers)
    for v in rows:
        ws.append([_priority_reason(v),v.get("cve"),v.get("title"),"YES" if v.get("kev") else "NO",v.get("epss_raw"),v.get("percentile_raw"),v.get("cvss_raw"),v.get("affected_assets"),v.get("findings"),v.get("published"),v.get("age_days"),v.get("nexpose_id")])
    _style_header(ws); _finish(ws); _table(ws,"T"+''.join(x for x in ws.title if x.isalnum())[:20]); _signal_colors(ws,headers)
    for colname in ("EPSS","EPSS Percentile"):
        ci=headers.index(colname)+1
        for r in range(2,ws.max_row+1): ws.cell(r,ci).number_format="0.00%"

def build_report(analysis):
    wb=Workbook(); wb.remove(wb.active)
    m=analysis["metrics"]; q=analysis["quality"]; status=analysis.get("intel_status",{})
    ws=wb.create_sheet("Executive Summary")
    ws.merge_cells("A1:F1"); ws["A1"]="Vulnerability Prioritization Report"; ws["A1"].fill=PatternFill("solid",fgColor=NAVY); ws["A1"].font=Font(color=WHITE,bold=True,size=18); ws["A1"].alignment=Alignment(vertical="center"); ws.row_dimensions[1].height=34
    ws["A3"]="Generated"; ws["B3"]=datetime.now().astimezone().strftime("%Y-%m-%d %I:%M %p %Z")
    sections=[("ENVIRONMENT",[("Assets Analyzed",m["assets"]),("Vulnerability Findings",m["findings"]),("Unique CVEs",m["cves"])]),("THREAT EXPOSURE",[("CISA KEV CVEs",m["kev_cves"]),("Assets with KEV",m["kev_assets"]),("High EPSS CVEs (≥10%)",m["high_epss_cves"]),("Critical Findings (CVSS ≥9)",m["critical_findings"])]),("SCAN FRESHNESS",[("Current (<3 days)",m["scan_current"]),("Aging (3–10 days)",m["scan_aging"]),("Stale (>10 days)",m["scan_stale"]),("Unknown",m["scan_unknown"])]),("DATA QUALITY",[("Missing CVE",q["missing_cve"]),("Missing Hostname",q["missing_hostname"]),("Missing CVSS v3",q["missing_cvss"]),("Missing Scan Date",q["missing_scan_date"])])]
    row=5
    for title,items in sections:
        ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=3); c=ws.cell(row,1,title); c.fill=PatternFill("solid",fgColor=BLUE); c.font=Font(color=WHITE,bold=True); row+=1
        for label,val in items: ws.cell(row,1,label); ws.cell(row,2,val); row+=1
        row+=1
    ws.cell(row,1,"COLOR LEGEND"); ws.cell(row,1).font=Font(bold=True); row+=1
    for text,color in [("Known exploited / Critical CVSS",RED),("High CVSS / very high EPSS",ORANGE),("Elevated EPSS / medium severity",YELLOW),("Lower technical severity",GREEN),("Missing / unknown data",GRAY)]:
        ws.cell(row,1,text); ws.cell(row,2," ").fill=PatternFill("solid",fgColor=color); row+=1
    row+=1; ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=6); ws.cell(row,1,"VulnPrioritizer does not assign a proprietary organizational risk score. Review priority is based on transparent CISA KEV, EPSS, CVSS, and exposure signals."); ws.cell(row,1).alignment=Alignment(wrap_text=True); ws.row_dimensions[row].height=34
    ws.column_dimensions["A"].width=38; ws.column_dimensions["B"].width=22; ws.column_dimensions["C"].width=18

    pri=[v for v in analysis["vulnerabilities"] if _priority_reason(v)]
    _write_vulns(wb.create_sheet("Priority Review"),pri)
    vws=wb.create_sheet("Vulnerability Overview"); _write_vulns(vws,analysis["vulnerabilities"])
    vws["A1"].comment=Comment("Default review order: CISA KEV first, then higher EPSS, higher CVSS, affected asset count, and vulnerability age as a final contextual tie-breaker. A KEV or high-EPSS CVE may rank above a CVSS 9+ CVE. This is not a proprietary risk score.","VulnPrioritizer")

    ws=wb.create_sheet("Asset Overview"); ah=["Asset ID","Hostname","IP Address","Last Scan","Scan Age (Days)","Scan Status","Priority Signals","Total Findings","Unique CVEs","CISA KEVs","EPSS ≥10%","CVSS ≥9 Findings","Maximum CVSS"]; ws.append(ah)
    for a in analysis["assets_table"]: ws.append([a["asset_id"],a["hostname"],a["ip_address"],a.get("last_scan_timestamp"),a.get("scan_age_days"),a.get("scan_status"),a.get("priority_signals"),a["findings"],a["unique_cves"],a["kev_cves"],a["high_epss_cves"],a["critical_findings"],float(a["max_cvss"]) if a["max_cvss"] else None])
    _style_header(ws); _finish(ws); _table(ws,"TAssetOverview")
    ws["A1"].comment=Comment("Default review order: CISA KEV CVE count, High EPSS CVE count (>=10%), Critical CVSS findings (>=9.0), maximum CVSS, then overall CVE/finding exposure. CMDB business context is not yet included. This is not a proprietary risk score.","VulnPrioritizer")
    for r in range(2,ws.max_row+1):
        scan_state=ws.cell(r,6).value
        if scan_state=="Stale": ws.cell(r,6).fill=PatternFill("solid",fgColor=RED); ws.cell(r,6).font=Font(color=RED_DARK,bold=True)
        elif scan_state=="Aging": ws.cell(r,6).fill=PatternFill("solid",fgColor=YELLOW)
        elif scan_state=="Current": ws.cell(r,6).fill=PatternFill("solid",fgColor=GREEN)
        if (ws.cell(r,10).value or 0)>0: ws.cell(r,10).fill=PatternFill("solid",fgColor=RED); ws.cell(r,10).font=Font(color=RED_DARK,bold=True)
        if (ws.cell(r,11).value or 0)>0: ws.cell(r,11).fill=PatternFill("solid",fgColor=ORANGE)
        if (ws.cell(r,12).value or 0)>0: ws.cell(r,12).fill=PatternFill("solid",fgColor=RED)

    kevrows=[v for v in analysis["vulnerabilities"] if v.get("kev")]
    ws=wb.create_sheet("CISA KEV"); kh=["CVE","Title","Vendor","Product","CVSS","EPSS","Affected Assets","Date Added to KEV","CISA Due Date","Known Ransomware Campaign Use","Required Action"]; ws.append(kh)
    for v in kevrows:
        k=analysis.get("kev_details",{}).get(v["cve"],{}); ws.append([v["cve"],v["title"],k.get("vendor"),k.get("product"),v.get("cvss_raw"),v.get("epss_raw"),v["affected_assets"],k.get("date_added"),k.get("due_date"),k.get("ransomware"),k.get("required_action")])
    _style_header(ws); _finish(ws); _table(ws,"TCISAKEV"); _signal_colors(ws,kh)
    for r in range(2,ws.max_row+1): ws.cell(r,6).number_format="0.00%"

    _write_vulns(wb.create_sheet("High EPSS"),[v for v in analysis["vulnerabilities"] if (v.get("epss_raw") or 0)>=.10])
    _write_vulns(wb.create_sheet("Critical CVSS"),[v for v in analysis["vulnerabilities"] if (v.get("cvss_raw") or 0)>=9])

    ws=wb.create_sheet("Data Quality"); ws.append(["Check","Count","Review Guidance"]); dq=[("Missing CVE",q["missing_cve"],"Cannot match to EPSS or CISA KEV without a CVE."),("Missing Hostname",q["missing_hostname"],"Review asset identification; Asset ID/IP may still be available."),("Missing CVSS v3",q["missing_cvss"],"Technical severity cannot be color-coded using CVSS v3."),("Missing Scan Date",q["missing_scan_date"],"Asset scan freshness cannot be determined.")]
    for x in dq: ws.append(x)
    _style_header(ws); _finish(ws); _table(ws,"TDataQuality")
    for r in range(2,ws.max_row+1):
        if (ws.cell(r,2).value or 0)>0: ws.cell(r,2).fill=PatternFill("solid",fgColor=YELLOW); ws.cell(r,2).font=Font(bold=True)

    ws=wb.create_sheet("Report Information"); ws.append(["Item","Value"]); info=[("Report purpose","Threat-enriched vulnerability review from a Rapid7 finding-level export."),("Priority model","No proprietary organizational risk score. Vulnerabilities: KEV -> EPSS -> CVSS -> affected assets -> age. Assets: KEV count -> High EPSS count -> Critical CVSS findings -> max CVSS -> overall exposure."),("High EPSS threshold","10% or greater."),("Critical CVSS threshold","9.0 or greater."),("Scan freshness","Current <3 days; Aging 3–10 days; Stale >10 days; Unknown = no valid scan date. Freshness is a data-confidence indicator and does not change priority ranking."),("Data handling","Rapid7 analysis and generated report files are temporary; public threat intelligence is stored separately."),("EPSS dataset",str(status.get("epss_dataset_date") or status.get("epss",{}).get("dataset_date") or "Unavailable")),("CISA KEV dataset",str(status.get("kev_dataset_date") or status.get("kev",{}).get("dataset_date") or "Unavailable"))]
    for x in info: ws.append(x)
    _style_header(ws); _finish(ws); _table(ws,"TReportInfo")

    ws=wb.create_sheet("Raw Rapid7 Data"); raw=analysis.get("raw_findings",[])
    rawheaders=["asset_id","ip_address","hostname","last_scan_date","nexpose_id","cve","title","date_published","severity_score","cvss_v3_score"]
    ws.append(rawheaders)
    for x in raw: ws.append([x.get(h) for h in rawheaders])
    _style_header(ws); _finish(ws); _table(ws,"TRawRapid7")

    out=BytesIO(); wb.save(out); out.seek(0); return out
