import ipaddress
from collections import Counter, defaultdict
from io import BytesIO
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

DEFAULT_PREFIX = 24
ALLOWED_PREFIXES = tuple(range(8, 33))

NETWORK_RULES = [
    {'network':'192.168.0.0/16','managed_by':'Not Corporate Owned','geolocation':''},
    {'network':'172.18.252.0/24','managed_by':'Service Express','geolocation':'UK'},
    {'network':'9.63.0.0/16','managed_by':'','geolocation':'NYC'},
]

def classify_network(ip_value):
    try:
        ip = ipaddress.ip_address(str(ip_value).strip())
        if ip.version != 4:
            return {'managed_by':'','geolocation':''}
    except ValueError:
        return {'managed_by':'','geolocation':''}
    for rule in NETWORK_RULES:
        if ip in ipaddress.ip_network(rule['network']):
            return {'managed_by':rule['managed_by'],'geolocation':rule['geolocation']}
    return {'managed_by':'','geolocation':''}

def normalize_prefix(value):
    try: p=int(value)
    except (TypeError, ValueError): return DEFAULT_PREFIX
    return p if p in ALLOWED_PREFIXES else DEFAULT_PREFIX

def build_network_breakdown(assets, prefix=DEFAULT_PREFIX):
    prefix=normalize_prefix(prefix)
    networks=defaultdict(list); invalid=[]
    for asset in assets:
        raw=str(asset.get('ip_address') or '').strip()
        try:
            ip=ipaddress.ip_address(raw)
            if ip.version != 4: raise ValueError('IPv4 only')
            net=str(ipaddress.ip_network(f'{ip}/{prefix}', strict=False))
            row=dict(asset); row['subnet']=net; row.update(classify_network(ip)); networks[net].append(row)
        except ValueError:
            row=dict(asset); row['subnet']='Unassigned / Invalid IP'; invalid.append(row)
    summary=[]
    for subnet, rows in sorted(networks.items(), key=lambda x: ipaddress.ip_network(x[0]).network_address):
        groups=Counter(r.get('asset_group','UNKNOWN') for r in rows)
        scans=Counter(r.get('scan_status','Unknown') for r in rows)
        ownership=Counter(r.get('managed_by','') for r in rows if r.get('managed_by','')); geos=Counter(r.get('geolocation','') for r in rows if r.get('geolocation','')); summary.append({'subnet':subnet,'asset_count':len(rows),'groups':dict(groups),'current':scans['Current'],'aging':scans['Aging'],'stale':scans['Stale'],'unknown_scan':scans['Unknown'],'managed_by':', '.join(sorted(ownership)),'geolocation':', '.join(sorted(geos))})
    return {'prefix':prefix,'subnet_count':len(summary),'asset_count':sum(x['asset_count'] for x in summary),'subnets':summary,'subnet_assets':dict(networks),'invalid_assets':invalid}

def _style_sheet(ws):
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:
        cell.font=Font(bold=True); cell.fill=PatternFill('solid', fgColor='D9EAF7'); cell.alignment=Alignment(wrap_text=True)
    for col in ws.columns:
        width=min(38,max(11,max(len(str(c.value or '')) for c in col)+2))
        ws.column_dimensions[get_column_letter(col[0].column)].width=width

def build_network_report(assets, prefix=DEFAULT_PREFIX):
    data=build_network_breakdown(assets,prefix)
    wb=Workbook(); ws=wb.active; ws.title='Network Summary'
    ws.append(['Observed Subnet','Unique Assets','Managed By','Geolocation','Current','Aging','Stale','Unknown Scan','Asset Group Breakdown'])
    for s in data['subnets']:
        breakdown=', '.join(f'{k}: {v}' for k,v in sorted(s['groups'].items()))
        ws.append([s['subnet'],s['asset_count'],s['managed_by'],s['geolocation'],s['current'],s['aging'],s['stale'],s['unknown_scan'],breakdown])
    _style_sheet(ws)
    detail=wb.create_sheet('Asset Detail')
    detail.append(['Observed Subnet','Network Managed By','Network Geolocation','Hostname','IP Address','Rapid7 Asset ID','Asset Group','Location','Management','Rapid7 OS','Last Scan','Scan Age (Days)','Scan Status'])
    for subnet in [s['subnet'] for s in data['subnets']]:
        for a in sorted(data['subnet_assets'][subnet],key=lambda x:str(x.get('ip_address',''))):
            detail.append([subnet,a.get('managed_by',''),a.get('geolocation',''),a.get('hostname'),a.get('ip_address'),a.get('asset_id'),a.get('asset_group'),a.get('location'),a.get('management'),a.get('operating_system'),a.get('last_scan_timestamp') or a.get('last_scan_date'),a.get('scan_age_days'),a.get('scan_status')])
    for a in data['invalid_assets']:
        detail.append(['Unassigned / Invalid IP','','',a.get('hostname'),a.get('ip_address'),a.get('asset_id'),a.get('asset_group'),a.get('location'),a.get('management'),a.get('operating_system'),a.get('last_scan_timestamp') or a.get('last_scan_date'),a.get('scan_age_days'),a.get('scan_status')])
    _style_sheet(detail)
    info=wb.create_sheet('Report Information')
    info.append(['Field','Value']); info.append(['Generated',datetime.now().strftime('%Y-%m-%d %H:%M:%S')]); info.append(['Subnet Prefix',f'/{data["prefix"]}']); info.append(['Observed Subnets',data['subnet_count']]); info.append(['Assets with valid IPv4',data['asset_count']]); info.append(['Unassigned / Invalid IP Assets',len(data['invalid_assets'])]); info.append(['Interpretation','Observed Subnets are inferred from asset IP addresses in this Rapid7 export. Presence indicates representation in the export; absence does not prove a subnet was not scanned.'])
    _style_sheet(info)
    out=BytesIO(); wb.save(out); out.seek(0); return out
