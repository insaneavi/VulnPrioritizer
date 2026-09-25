from io import BytesIO
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

STATUS_COLORS={"Current":"DCFCE7","Aging":"FEF3C7","Stale":"FECACA","Unknown":"E5E7EB"}
STATUS_FONT={"Current":"14532D","Aging":"713F12","Stale":"7F1D1D","Unknown":"374151"}
NAVY="17233D"; WHITE="FFFFFF"; BORDER="D7DEE9"; RED="7F1D1D"; RED_FILL="FECACA"; MUTED="5B6577"

def _border():
    side=Side(style="thin", color=BORDER)
    return Border(left=side,right=side,top=side,bottom=side)

def _header(cell):
    cell.font=Font(bold=True,color=WHITE)
    cell.fill=PatternFill("solid",fgColor=NAVY)
    cell.border=_border()

def _body(cell):
    cell.border=_border()

def _unknown(cell):
    cell.font=Font(bold=True,color=RED)
    cell.fill=PatternFill("solid",fgColor=RED_FILL)
    cell.border=_border()

def build_asset_classification_report(analysis):
    output=BytesIO()
    wb=Workbook()
    wb.remove(wb.active)
    generated=datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    assets=analysis.get("assets_table",[])

    summary=wb.create_sheet("Classification Summary")
    summary.merge_cells("A1:F1")
    summary["A1"]="VulnPrioritizer Asset Classification Review"
    summary["A1"].font=Font(bold=True,size=18,color=WHITE)
    summary["A1"].fill=PatternFill("solid",fgColor=NAVY)
    summary["A1"].alignment=Alignment(vertical="center")
    summary.row_dimensions[1].height=28
    for col,width in {"A":30,"B":18,"D":28,"E":28,"F":28}.items(): summary.column_dimensions[col].width=width
    summary["A3"]="Generated"; _header(summary["A3"]); summary["B3"]=generated; _body(summary["B3"])
    summary["A4"]="Total Assets"; _header(summary["A4"]); summary["B4"]=len(assets); _body(summary["B4"])
    unknown_count=sum(1 for a in assets if a.get("asset_group")=="UNKNOWN")
    summary["A5"]="UNKNOWN Assets"; _header(summary["A5"]); summary["B5"]=unknown_count; (_unknown if unknown_count else _body)(summary["B5"])
    summary["A7"]="Asset Group"; _header(summary["A7"]); summary["B7"]="Count"; _header(summary["B7"])
    counts={}
    for a in assets:
        group=a.get("asset_group","UNKNOWN")
        counts[group]=counts.get(group,0)+1
    for row,(group,count) in enumerate(sorted(counts.items()),8):
        summary.cell(row,1,group); (_unknown if group=="UNKNOWN" else _body)(summary.cell(row,1))
        summary.cell(row,2,count); _body(summary.cell(row,2))
    summary["D3"]="Purpose"; _header(summary["D3"])
    summary.merge_cells("D4:F6")
    summary["D4"]="Audit the operational asset tags derived from Rapid7 hostname and OS data. UNKNOWN assets are intentionally exposed for research rather than guessed."
    summary["D4"].font=Font(color=MUTED,italic=True)
    summary["D4"].alignment=Alignment(wrap_text=True,vertical="top")

    headers=["Hostname","IP Address","Rapid7 Asset ID","Asset Group","Location","Management","Classification Source","Classification Rule","Rapid7 Operating System","Last Scan","Scan Age (Days)","Scan Status"]
    widths=[22,16,18,22,14,16,22,30,36,21,15,14]
    def detail_sheet(name,rows):
        ws=wb.create_sheet(name)
        ws.freeze_panes="A2"
        for c,(h,w) in enumerate(zip(headers,widths),1):
            ws.cell(1,c,h); _header(ws.cell(1,c)); ws.column_dimensions[get_column_letter(c)].width=w
        for r,a in enumerate(rows,2):
            values=[a.get("hostname"),a.get("ip_address"),a.get("asset_id"),a.get("asset_group"),a.get("location"),a.get("management"),a.get("classification_source"),a.get("classification_rule"),a.get("operating_system"),a.get("last_scan_timestamp") or "",a.get("scan_age_days"),a.get("scan_status")]
            for c,v in enumerate(values,1):
                cell=ws.cell(r,c,"" if v is None else v); _body(cell)
                if c==4 and a.get("asset_group")=="UNKNOWN": _unknown(cell)
                if c==12:
                    status=a.get("scan_status","Unknown")
                    cell.font=Font(bold=True,color=STATUS_FONT.get(status,"374151"))
                    cell.fill=PatternFill("solid",fgColor=STATUS_COLORS.get(status,"E5E7EB"))
        last=max(2,len(rows)+1)
        ws.auto_filter.ref=f"A1:L{last}"
    detail_sheet("Classification Detail",assets)
    detail_sheet("UNKNOWN Assets",[a for a in assets if a.get("asset_group")=="UNKNOWN"])
    wb.save(output); output.seek(0); return output
