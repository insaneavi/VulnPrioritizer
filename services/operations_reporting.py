from io import BytesIO
from datetime import datetime
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY="172033"; BLUE="2563EB"; WHITE="FFFFFF"; RED="FECACA"; RED_DARK="991B1B"
ORANGE="FED7AA"; YELLOW="FEF3C7"; GREEN="DCFCE7"; GRAY="E5E7EB"; LIGHT="F8FAFC"
THIN=Side(style="thin",color="D1D5DB")

def _generated():
    return datetime.now().astimezone().strftime("%B %d, %Y %I:%M %p %Z")

def _os_group(row):
    name=(row.get("os_name") or "").lower()
    desc=(row.get("operating_system") or "").lower()
    text=f"{name} {desc}"
    if "windows 11" in text: return "Windows 11"
    if "windows 10" in text: return "Windows 10"
    m=re.search(r"windows server\s+(2008 r2|2008|2012 r2|2012|2016|2019|2022|2025)",text)
    if m: return "Windows Server "+m.group(1).title().replace("R2","R2")
    if "windows server" in text: return "Windows Server - Version Unknown"
    if "windows" in text: return "Windows - Version Unknown"
    return ""

def _is_windows_finding(r):
    title=(r.get("title") or "").lower()
    return title.startswith("microsoft windows:") or "windows security update" in title or "windows cumulative update" in title

def _is_edge_finding(r):
    title=(r.get("title") or "").lower()
    return "microsoft edge" in title or "edge (chromium" in title

def _kev_map(analysis):
    return {v["cve"]:bool(v.get("kev")) for v in analysis.get("vulnerabilities",[])}

def _epss_map(analysis):
    return {v["cve"]:float(v.get("epss_raw") or 0) for v in analysis.get("vulnerabilities",[])}

def _campaigns(analysis):
    raw=analysis.get("raw_findings",[])
    asset_index=analysis.get("asset_index",{})
    kev=_kev_map(analysis); epss=_epss_map(analysis)
    campaigns=[]

    # Windows monthly-update campaigns are based on OS inventory + Windows-specific findings.
    groups={}
    for r in raw:
        aid=str(r.get("asset_id") or "")
        a=asset_index.get(aid,{})
        g=_os_group(a if a else r)
        if g and _is_windows_finding(r):
            groups.setdefault(g,[]).append(r)
    order=["Windows 11","Windows 10","Windows Server 2025","Windows Server 2022","Windows Server 2019","Windows Server 2016","Windows Server 2012 R2","Windows Server 2012","Windows Server 2008 R2","Windows Server 2008","Windows Server - Version Unknown","Windows - Version Unknown"]
    for g in sorted(groups,key=lambda x:(order.index(x) if x in order else 99,x)):
        rows=groups[g]
        campaigns.append(_make_campaign(
            g+" Monthly Patching",g,
            "Apply the latest organization-approved Microsoft monthly cumulative/security updates for this operating-system group. Validate deployment through the standard patch-management process and rescan afterward.",
            rows,asset_index,kev,epss))

    # Edge is a separate software campaign because updating Edge can address many CVEs at once.
    edge=[r for r in raw if _is_edge_finding(r)]
    if edge:
        campaigns.append(_make_campaign(
            "Microsoft Edge Update","Microsoft Edge",
            "Update Microsoft Edge to the latest organization-approved enterprise release using the standard software deployment mechanism, then validate the installed version and rescan.",
            edge,asset_index,kev,epss))
    return campaigns

def _make_campaign(name,category,action,rows,asset_index,kev,epss):
    aids=sorted({str(r.get("asset_id") or "") for r in rows if r.get("asset_id") not in (None,"")})
    cves={str(r.get("cve") or "").upper() for r in rows if str(r.get("cve") or "").upper().startswith("CVE-")}
    assets=[]
    for aid in aids:
        a=asset_index.get(aid,{})
        ar=[r for r in rows if str(r.get("asset_id") or "")==aid]
        acves={str(r.get("cve") or "").upper() for r in ar if str(r.get("cve") or "").upper().startswith("CVE-")}
        assets.append({
            "asset_id":aid,"hostname":a.get("hostname") or next((str(x.get("hostname") or "") for x in ar if x.get("hostname")),"—"),
            "ip_address":a.get("ip_address") or next((str(x.get("ip_address") or "") for x in ar if x.get("ip_address")),"—"),
            "operating_system":a.get("operating_system") or next((str(x.get("operating_system") or "") for x in ar if x.get("operating_system")),"Unknown"),
            "last_scan_date":a.get("last_scan_timestamp") or "Unknown",
            "scan_age_days":a.get("scan_age_days"),
            "scan_status":a.get("scan_status") or "Unknown",
            "findings":len(ar),"unique_cves":len(acves),
            "kev_cves":sum(1 for c in acves if kev.get(c)),
            "high_epss_cves":sum(1 for c in acves if epss.get(c,0)>=.10),
            "critical_findings":sum(1 for x in ar if _num(x.get("cvss_v3_score"))>=9)
        })
    return {"name":name,"category":category,"action":action,"rows":rows,"assets":assets,
            "asset_count":len(aids),"findings":len(rows),"unique_cves":len(cves),
            "kev_cves":sum(1 for c in cves if kev.get(c)),
            "high_epss_cves":sum(1 for c in cves if epss.get(c,0)>=.10)}

def _num(v):
    try: return float(v)
    except Exception: return 0

def _header(ws,row):
    for c in ws[row]:
        c.fill=PatternFill("solid",fgColor=NAVY); c.font=Font(color=WHITE,bold=True); c.alignment=Alignment(vertical="center")
    ws.row_dimensions[row].height=24

def _widths(ws):
    for col in range(1,ws.max_column+1):
        vals=[str(ws.cell(r,col).value or "") for r in range(1,min(ws.max_row,300)+1)]
        ws.column_dimensions[get_column_letter(col)].width=min(max(11,max(map(len,vals),default=8)+2),48)
    for row in ws.iter_rows():
        for c in row:
            c.border=Border(bottom=THIN); c.alignment=Alignment(vertical="top",wrap_text=True)

def _table(ws,start,end,name):
    if end<=start: return
    ref=f"A{start}:{get_column_letter(ws.max_column)}{end}"
    t=Table(displayName=name[:30],ref=ref)
    t.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True)
    ws.add_table(t)

def _campaign_sheet(ws,c,index):
    ws.merge_cells("A1:L1"); ws["A1"]="PATCH CAMPAIGN — "+c["name"]; ws["A1"].fill=PatternFill("solid",fgColor=NAVY); ws["A1"].font=Font(color=WHITE,bold=True,size=16); ws.row_dimensions[1].height=32
    ws["A3"]="Generated"; ws["B3"]=_generated()
    metrics=[("Assets Requiring Attention",c["asset_count"]),("Vulnerability Findings",c["findings"]),("Unique CVEs",c["unique_cves"]),("CISA KEV CVEs",c["kev_cves"]),("High EPSS CVEs (≥10%)",c["high_epss_cves"])]
    r=5
    for k,v in metrics: ws.cell(r,1,k).font=Font(bold=True); ws.cell(r,2,v); r+=1
    ws.cell(r+1,1,"Recommended Action").font=Font(bold=True,color=WHITE); ws.cell(r+1,1).fill=PatternFill("solid",fgColor=BLUE)
    ws.merge_cells(start_row=r+2,start_column=1,end_row=r+3,end_column=12); ws.cell(r+2,1,c["action"]); ws.cell(r+2,1).alignment=Alignment(wrap_text=True,vertical="top")
    h=r+5; headers=["Hostname","IP Address","Asset ID","Operating System","Last Scan","Scan Age (Days)","Scan Status","Findings Addressed","Unique CVEs","CISA KEVs","High EPSS CVEs","Critical Findings"]; 
    for i,x in enumerate(headers,1): ws.cell(h,i,x)
    _header(ws,h)
    for a in c["assets"]:
        ws.append([a["hostname"],a["ip_address"],a["asset_id"],a["operating_system"],a["last_scan_date"],a["scan_age_days"],a["scan_status"],a["findings"],a["unique_cves"],a["kev_cves"],a["high_epss_cves"],a["critical_findings"]])
    _table(ws,h,ws.max_row,"CampaignAssets"+str(index))
    ws.freeze_panes=f"A{h+1}"
    for rr in range(h+1,ws.max_row+1):
        status=ws.cell(rr,7).value
        if status=="Stale": ws.cell(rr,7).fill=PatternFill("solid",fgColor=RED); ws.cell(rr,7).font=Font(color=RED_DARK,bold=True)
        elif status=="Aging": ws.cell(rr,7).fill=PatternFill("solid",fgColor=YELLOW)
        elif status=="Current": ws.cell(rr,7).fill=PatternFill("solid",fgColor=GREEN)
        if (ws.cell(rr,10).value or 0)>0: ws.cell(rr,10).fill=PatternFill("solid",fgColor=RED); ws.cell(rr,10).font=Font(color=RED_DARK,bold=True)
        if (ws.cell(rr,11).value or 0)>0: ws.cell(rr,11).fill=PatternFill("solid",fgColor=ORANGE)
        if (ws.cell(rr,12).value or 0)>0: ws.cell(rr,12).fill=PatternFill("solid",fgColor=RED)
    _widths(ws)

def build_operations_report(analysis):
    campaigns=_campaigns(analysis)
    wb=Workbook(); ws=wb.active; ws.title="Operations Summary"
    ws.merge_cells("A1:H1"); ws["A1"]="Operations Patching Report"; ws["A1"].fill=PatternFill("solid",fgColor=NAVY); ws["A1"].font=Font(color=WHITE,bold=True,size=18); ws.row_dimensions[1].height=34
    ws["A3"]="Generated"; ws["B3"]=_generated()
    ws["A4"]="Purpose"; ws["B4"]="Translate Rapid7 vulnerability findings into downstream remediation campaigns for Operations. Campaigns consolidate many findings into practical patching actions."
    ws["A6"]="Important"; ws["B6"]="Recommended actions are grouping guidance, not proof that one update remediates every listed finding. Operations should follow approved patch/change procedures and validate with a post-patch scan."
    ws["A7"]="Scan Freshness"; ws["B7"]="Current <3 days; Aging 3–10 days; Stale >10 days; Unknown = no valid scan date. Freshness indicates confidence in the vulnerability evidence and does not change priority ranking."
    headers=["Patch Campaign","Assets","Findings Addressed","Unique CVEs","CISA KEVs","High EPSS CVEs","Recommended Action"]
    for i,h in enumerate(headers,1): ws.cell(8,i,h)
    _header(ws,8)
    for c in campaigns: ws.append([c["name"],c["asset_count"],c["findings"],c["unique_cves"],c["kev_cves"],c["high_epss_cves"],c["action"]])
    _table(ws,8,ws.max_row,"OperationsCampaigns"); ws.freeze_panes="A9"; _widths(ws)
    for r in range(9,ws.max_row+1):
        if (ws.cell(r,5).value or 0)>0: ws.cell(r,5).fill=PatternFill("solid",fgColor=RED); ws.cell(r,5).font=Font(color=RED_DARK,bold=True)
        if (ws.cell(r,6).value or 0)>0: ws.cell(r,6).fill=PatternFill("solid",fgColor=ORANGE)

    # One technical worksheet per campaign, with generation timestamp on every campaign.
    used=set()
    for i,c in enumerate(campaigns,1):
        base=re.sub(r'[\[\]:*?/\\\\]',' ',c["category"])[:28].strip() or f"Campaign {i}"
        name=base; n=2
        while name in used: name=(base[:25]+f" {n}"); n+=1
        used.add(name)
        _campaign_sheet(wb.create_sheet(name),c,i)

    # Consolidated finding-level campaign data for filtering / ticket creation.
    ws=wb.create_sheet("Campaign Finding Data")
    headers=["Patch Campaign","Asset ID","Hostname","IP Address","Operating System","Last Scan","Scan Age (Days)","Scan Status","CVE","Title","CVSS","Rapid7 Vulnerability ID"]
    ws.append(headers); _header(ws,1)
    for c in campaigns:
        for r in c["rows"]:
            aid=str(r.get("asset_id") or ""); a=analysis.get("asset_index",{}).get(aid,{})
            ws.append([c["name"],aid,a.get("hostname") or r.get("hostname"),a.get("ip_address") or r.get("ip_address"),a.get("operating_system") or r.get("operating_system"),a.get("last_scan_timestamp") or "Unknown",a.get("scan_age_days"),a.get("scan_status") or "Unknown",r.get("cve"),r.get("title"),r.get("cvss_v3_score"),r.get("nexpose_id")])
    _table(ws,1,ws.max_row,"CampaignFindingData"); ws.freeze_panes="A2"; _widths(ws)

    out=BytesIO(); wb.save(out); out.seek(0); return out, campaigns
